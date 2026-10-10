import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_document():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Arial'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(4)

    def set_cell_background(cell, fill_hex):
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shd)

    def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def add_heading_clean(text, level=1):
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Arial'
        if level == 1:
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
        elif level == 2:
            run.font.size = Pt(11.5)
            run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
        return p

    def add_bullet(text, bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.bold = True
            r_bold.font.name = 'Arial'
            r_bold.font.size = Pt(10.5)
            r_bold.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
        r_text = p.add_run(text)
        r_text.font.name = 'Arial'
        r_text.font.size = Pt(10.5)
        r_text.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        return p

    def add_code_block(code_text):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.1
        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x2A, 0x2A, 0x2A)
        return p

    # --- ENCABEZADO PRINCIPAL ---
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_after = Pt(2)
    p_title.paragraph_format.space_before = Pt(0)
    r = p_title.add_run("Documento de Requerimientos y Pruebas")
    r.bold = True
    r.font.name = 'Arial'
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(0x11, 0x11, 0x11)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(4)
    r = p_sub.add_run("Pruebas Unitarias e Integración — Usuarios API (Node.js + Jest)")
    r.bold = True
    r.font.name = 'Arial'
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    p_proj = doc.add_paragraph()
    p_proj.paragraph_format.space_after = Pt(2)
    r = p_proj.add_run("Arquitectura réplica de WebApp adaptada al stack JavaScript / Node.js")
    r.font.name = 'Arial'
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(8)
    r = p_meta.add_run("Proyecto base: Express.js (v5) + Sequelize ORM + SQLite3 + bcryptjs\n"
                       "Framework de pruebas: Jest + Supertest (31 pruebas automatizadas, 100% aprobadas)\n"
                       "Contenedor: Docker (Node 22 Alpine) | Puerto HTTP: 3000")
    r.font.name = 'Arial'
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    # --- 1. OBJETIVO ---
    add_heading_clean("1. Objetivo", 1)
    doc.add_paragraph(
        "Replicar rigurosamente la arquitectura, estándares de calidad y exhaustividad de pruebas del proyecto WebApp (.NET) "
        "en la API de administración de usuarios y roles (usuarios-api), empleando las tecnologías nativas del ecosistema JavaScript "
        "(Node.js, Express, Sequelize, SQLite) y el framework de pruebas automatizadas Jest junto con Supertest. La suite implementada "
        "cubre un total de 31 casos de prueba automatizados, evaluando tanto operaciones exitosas (Happy Path) como escenarios de error "
        "de cliente (nombres vacíos, llaves foráneas inválidas, identificadores inexistentes, conflictos de unicidad y protección relacional)."
    )

    # --- 2. ALCANCE ---
    add_heading_clean("2. Alcance", 1)
    add_bullet("Implementación de los métodos PUT faltantes en el backend (/api/roles/:id y /api/usuarios/:id), completando los 4 verbos HTTP (GET, POST, PUT, DELETE) para todas las entidades.")
    add_bullet("Modularización de la aplicación separando la definición de la app Express (app.js) del proceso de arranque HTTP (server.js) para permitir ejecución aislada con Supertest sin ocupar puertos de red.")
    add_bullet("Configuración de entorno de pruebas aislado (NODE_ENV=test) con SQLite dedicado (database.test.sqlite), asegurando que las pruebas no alteren la base de desarrollo ni colisionen en paralelo.")
    add_bullet("Pruebas de mantenimiento de base de datos con respaldo descargable (/api/database/backup) y vaciado transaccional seguro (/api/database/vaciar).")

    # --- 3. TECNOLOGÍA Y HERRAMIENTAS ---
    add_heading_clean("3. Tecnologías y herramientas utilizadas", 1)
    add_bullet("Servidor HTTP y enrutamiento minimalista y eficiente.", "Express 5.1: ")
    add_bullet("Mapeo objeto-relacional para gestión de esquemas, relaciones 1:N entre Roles y Usuarios y transacciones.", "Sequelize 6.37 + SQLite3: ")
    add_bullet("Cifrado unidireccional de contraseñas de usuarios con salt de 10 rondas.", "bcryptjs: ")
    add_bullet("Framework de pruebas unitarias y aserciones en JavaScript (equivalente a xUnit en .NET).", "Jest 29.7: ")
    add_bullet("Cliente HTTP programático para simular peticiones reales contra Express sin necesidad de levantar servidores en red.", "Supertest 7.0: ")

    # --- 4. MATRIZ DE ENDPOINTS ---
    add_heading_clean("4. Matriz completa de endpoints", 1)
    table0 = doc.add_table(rows=1, cols=4)
    table0.alignment = WD_TABLE_ALIGNMENT.CENTER
    table0.autofit = False

    hdr_cells = table0.rows[0].cells
    hdr_titles = ['#', 'Endpoint', 'Método', 'Descripción']
    col_widths = [Inches(0.4), Inches(2.2), Inches(0.9), Inches(3.0)]

    for idx, name in enumerate(hdr_titles):
        cell = hdr_cells[idx]
        cell.text = name
        cell.width = col_widths[idx]
        set_cell_background(cell, "EAEAEA")
        set_cell_margins(cell, top=120, bottom=120, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in [0, 2] else WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            r.bold = True
            r.font.name = 'Arial'
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(0x11, 0x11, 0x11)

    endpoints_data = [
        ('1', '/', 'GET', 'Comprobación de estado y salud del servidor.'),
        ('2', '/api/roles', 'GET', 'Lista todos los roles de usuario registrados.'),
        ('3', '/api/roles/:id', 'GET', 'Obtiene un rol por su ID con sus usuarios asociados.'),
        ('4', '/api/roles', 'POST', 'Crea un nuevo rol validando unicidad y no vacío.'),
        ('5', '/api/roles/:id', 'PUT', 'Actualiza el nombre de un rol existente.'),
        ('6', '/api/roles/:id', 'DELETE', 'Elimina un rol validando que no tenga usuarios dependientes.'),
        ('7', '/api/usuarios', 'GET', 'Lista todos los usuarios (excluye hashes de contraseñas).'),
        ('8', '/api/usuarios/:id', 'GET', 'Obtiene el detalle de un usuario por su ID.'),
        ('9', '/api/usuarios', 'POST', 'Registra un nuevo usuario cifrando su contraseña y validando rolId.'),
        ('10', '/api/usuarios/:id', 'PUT', 'Actualiza datos de un usuario (nombre, correo, rolId, password).'),
        ('11', '/api/usuarios/:id', 'DELETE', 'Elimina permanentemente un usuario por su ID.'),
        ('12', '/api/database/backup', 'POST', 'Crea una copia de la base SQLite y la envía como descarga.'),
        ('13', '/api/database/vaciar', 'DELETE / POST', 'Vacía controladamente usuarios y roles mediante transacción.')
    ]

    for row_data in endpoints_data:
        row = table0.add_row()
        for idx, text in enumerate(row_data):
            cell = row.cells[idx]
            cell.text = text
            cell.width = col_widths[idx]
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.1
            p.paragraph_format.space_after = Pt(2)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in [0, 2] else WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = 'Arial'
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
                if idx == 2:
                    r.bold = True

    # --- 5. ESTRUCTURA DE LA SUITE DE PRUEBAS JEST ---
    add_heading_clean("5. Estructura de la suite de pruebas automatizadas (Jest + Supertest)", 1)
    doc.add_paragraph(
        "Las pruebas se organizaron en el directorio tests/ cubriendo 3 suites modulares con un total de 31 casos de prueba ejecutados satisfactoriamente:"
    )
    add_code_block(
        "usuarios-api/tests/\n"
        "├── setup.js          # Sincronización limpia de SQLite y cierre de conexión\n"
        "├── roles.test.js     # 13 pruebas: CRUD, validación de vacíos, 404, unicidad (409) y protección relacional\n"
        "├── usuarios.test.js  # 14 pruebas: CRUD, validación de correo/password, 404, llave foránea rolId (400) y duplicados\n"
        "└── database.test.js  # 4 pruebas: vaciado transaccional, método alternativo POST, resiliencia y backup\n"
        "Total: 31 pruebas unitarias e integración (100% aprobadas)."
    )

    output_path = "Requerimientos_Pruebas_Unitarias_Usuarios_API.docx"
    doc.save(output_path)
    print(f"Documento guardado con éxito en: {output_path}")

if __name__ == "__main__":
    create_document()
